import sys, re, datetime as dt, pandas as pd
from openpyxl import load_workbook
f = sys.argv[1]; wb = load_workbook(f)
err = []
print('시트:', wb.sheetnames)
if wb.sheetnames != ['장소_대응표', '제보']: err.append('시트 구성이 양식과 다름')
pl = pd.read_excel(f, sheet_name='장소_대응표', dtype=str).dropna(subset=['zone_id'])
rp = pd.read_excel(f, sheet_name='제보', dtype=str).fillna('')
T1 = ['place_text', 'zone_id', '위치 설명']; T2 = ['report_id', 'participant_id', 'raw_text', 'sent_at', 'observed_date', 'observed_time', 'time_precision', 'place_text', 'zone_id', 'intensity', 'note']
if list(pl.columns[:3]) != T1: err.append('장소_대응표 앞 3열이 양식과 다름')
if list(rp.columns[:11]) != T2: err.append('제보 앞 11열이 양식과 다름')
zmap = dict(zip(pl.place_text, pl.zone_id))
if set(pl.zone_id) - {'U01', 'U02', 'U03', 'H01', 'OUT', 'G01'}: err.append('zone_id 허용값 밖')
PART = {'04:00', '07:30', '14:00', '19:00', '23:00'}
ids = rp.report_id.tolist()
if ids != [f'R{i:04d}' for i in range(1, len(ids) + 1)]: err.append('report_id가 R0001부터 연속이 아님')
for r in rp.itertuples():
    i = r.report_id; inc = r.분석_포함 == 'Y'
    if r.participant_id not in {'P00', 'P01', 'P02', 'P03', 'P04', 'P05'}: err.append(f'{i} participant_id')
    if r.sent_at and not re.fullmatch(r'\d{4}-\d{2}-\d{2} \d{2}:\d{2}', r.sent_at): err.append(f'{i} sent_at 형식')
    if r.time_precision not in {'exact', 'hour', 'part', 'unknown'}: err.append(f'{i} time_precision')
    if r.observed_date and not re.fullmatch(r'\d{4}-\d{2}-\d{2}', r.observed_date): err.append(f'{i} observed_date 형식')
    if r.observed_time and not re.fullmatch(r'\d{2}:\d{2}', r.observed_time): err.append(f'{i} observed_time 형식')
    if r.intensity and r.intensity not in {'1', '2', '3'}: err.append(f'{i} intensity')
    if not inc:
        if not r.note: err.append(f'{i} 제외 줄인데 note에 이유 없음')
        continue
    if not r.observed_date: err.append(f'{i} observed_date 비어 있음')
    if r.time_precision == 'unknown' and r.observed_time: err.append(f'{i} unknown인데 시각이 있음')
    if r.time_precision != 'unknown' and not r.observed_time: err.append(f'{i} 시각이 비어 있음')
    if r.time_precision == 'part' and r.observed_time not in PART: err.append(f'{i} part 대표 시각이 아님 {r.observed_time}')
    if r.place_text not in zmap: err.append(f'{i} place_text가 장소_대응표에 없음 "{r.place_text}"')
    elif zmap[r.place_text] != r.zone_id: err.append(f'{i} zone_id가 장소_대응표와 다름')
    # 본문 시각 우선, 없으면 보낸 시각
    t = r.raw_text; sa = dt.datetime.fromisoformat(r.sent_at) if r.sent_at else None
    has_time = re.search(r'\d{1,2}:\d{2}|\d+시|새벽|아침|낮|저녁|밤|\d+/\d+', t)
    obs = dt.datetime.fromisoformat(f'{r.observed_date} {r.observed_time or "00:00"}')
    if not has_time and r.time_precision in ('hour', 'exact') and sa and obs != sa: err.append(f'{i} 본문 시각 없는데 보낸 시각과 다름')
    if re.search(r'어젯?\s?밤|어제\s?밤', t) and (obs.date() != sa.date() - dt.timedelta(days=1) or r.observed_time != '23:00'): err.append(f'{i} 어젯밤 규칙')
    if '새벽' in t and r.time_precision == 'part' and r.observed_time != '04:00': err.append(f'{i} 새벽 규칙')
    if obs > (sa + dt.timedelta(minutes=1) if sa else obs): err.append(f'{i} 관측이 전송보다 늦음')
    m = re.search(r'(?<![/\d:])([123])(?![/\d:])', t.replace('[이름]', ''))
    if not r.intensity and m and r.time_precision != 'unknown' and not re.search(r'\d+분째|\d+/\d+', t): err.append(f'{i} 원문에 숫자가 있는데 intensity 비어 있음')
print('제보', len(rp), '줄 / 분석 포함', (rp.분석_포함 == 'Y').sum(), '/ 제외', (rp.분석_포함 == 'N').sum())
print('장소', len(pl), '곳, 쓰인 구역', sorted(set(pl.zone_id)))
print('오류', len(err)); [print(' ', e) for e in err[:40]]
print(rp[rp.분석_포함 == 'N'][['report_id', 'raw_text', 'note']].to_string())
