"""reports_raw.csv + reports.csv + 장소_대응표.csv → 냄새나침반_주민제보_입력표.xlsx (시트 2개)"""
import sys, os, re, datetime as dt
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

D = sys.argv[1]; OUTF = sys.argv[2]
raw = pd.read_csv(D + 'reports_raw.csv', dtype=str).fillna('')
rep = pd.read_csv(D + 'reports.csv', dtype=str).fillna('')
pl = pd.read_csv(D + '장소_대응표.csv', dtype=str).fillna('')
repm = {r.report_id: r for r in rep.itertuples()}

wb = Workbook()
HDR = PatternFill('solid', fgColor='DDE7F3'); EXT = PatternFill('solid', fgColor='EFEFEF')
def header(ws, cols, ext_from=None, notes=None):
    ws.append(cols)
    for i, c in enumerate(ws[1], 1):
        c.font = Font(bold=True); c.fill = EXT if ext_from and i >= ext_from else HDR
        c.alignment = Alignment(vertical='center')
        if notes and cols[i - 1] in notes: c.comment = Comment(notes[cols[i - 1]], 'lyn')
    ws.freeze_panes = 'A2'

# ① 장소_대응표
ws1 = wb.active; ws1.title = '장소_대응표'
header(ws1, ['place_text', 'zone_id', '위치 설명', '톡 표현', 'lat', 'lon'], ext_from=4)
for r in pl.itertuples():
    desc = r.실제_장소 if r.실제_장소 != r.place_text else r.place_text
    if r.zone_id == 'OUT': desc += ' (운양동 밖)'
    if r.zone_id == 'G01': desc += ' (대조 지역)'
    ws1.append([r.place_text, r.zone_id, desc, r.톡_표현, float(r.lat), float(r.lon)])
dv1 = DataValidation(type='list', formula1='"U01,U02,U03,G01,H01,OUT"', allow_blank=False)
ws1.add_data_validation(dv1); dv1.add(f'B2:B{ws1.max_row + 50}')
ws1.append([]); ws1.append(['zone_id: U01 한강변 북측 / U02 중앙(운양역) / U03 동측·모담산 / G01 고촌(대조, 추가) / H01 하성(이번 자료에는 없음) / OUT 운양동 밖'])

# ② 제보
ws2 = wb.create_sheet('제보')
BASE = ['report_id', 'participant_id', 'raw_text', 'sent_at', 'observed_date', 'observed_time', 'time_precision', 'place_text', 'zone_id', 'intensity', 'note']
EXTC = ['분석_포함', 'source', 'source_ref', 'time_basis', 'window_start', 'window_end', 'place_basis', 'source_check']
notes = {'분석_포함': 'Y = reports.csv에 들어간 줄. N = 날짜 없음·기간 밖·중복 병합으로 빠진 줄(note에 이유)',
         'time_basis': '시각을 어디서 가져왔는지. 확장열_설명.md 참고', 'window_start': '원문으로 볼 때 냄새를 맡았을 수 있는 구간 시작',
         'place_basis': 'stated = 원문에 장소 있음 / assumed_home = 장소가 없어 집으로 가정', 'source_check': '원문 대조 가능 여부'}
header(ws2, BASE + EXTC, ext_from=len(BASE) + 1, notes=notes)

def excl_reason(r):
    if not r.sent_at: return '날짜 없음(캡처에 날짜 구분선 없음) → 분석 제외'
    t = r.raw_time_text
    return None

for r in raw.itertuples():
    q = repm.get(r.raw_id)
    if q is not None:
        tm = '' if q.time_precision == 'unknown' else q.observed_time
        note = q.note
        if q.time_precision == 'unknown': note = '; '.join([n for n in note.split('; ') if n] + ['시각 모름 → observed_time 비움'])
        ws2.append([r.raw_id, q.participant_id, r.raw_text, r.sent_at, q.observed_date, tm, q.time_precision, q.place_text, q.zone_id,
                    int(q.intensity) if q.intensity else None, note, 'Y', r.source, r.source_ref, q.time_basis, q.window_start, q.window_end,
                    q.place_basis, q.source_check])
    else:
        # 정리본에서 빠진 줄: 이유를 note에
        why = ''; od = ''; ot = ''; prec = ''; place = ''; zone = ''; inten = None
        dup = rep[rep.note.str.contains(f'중복 병합 {r.raw_id}') | rep.note.str.contains(f'인스타 {r.raw_id} 제외')]
        if not r.sent_at:
            why = '날짜 없음(캡처에 날짜 구분선 없음) → 분석 제외'; prec = 'unknown'
        elif len(dup):
            why = f'{dup.iloc[0].report_id} 줄과 같은 제보 → 병합, 분석 제외'
        else:
            sa = dt.datetime.fromisoformat(r.sent_at)
            if r.raw_time_text == '지난주 토요일':
                d0 = sa.date() - dt.timedelta(days=(sa.weekday() - 5) % 7 + 7 if sa.weekday() != 5 else 7)
                od = d0.isoformat(); prec = 'unknown'
            elif r.raw_time_text == '오늘 아침':
                od = sa.date().isoformat(); ot = '07:30'; prec = 'part'
            else:
                od = sa.date().isoformat(); ot = sa.strftime('%H:%M'); prec = 'hour'
            why = '기간(2026-06-20~09-30) 밖 → 분석 제외'
        if len(dup):
            k_ = dup.iloc[0]; od, ot, prec, place, zone = k_.observed_date, k_.observed_time, k_.time_precision, k_.place_text, k_.zone_id
        else:
            home = {'P01': ['한라'], 'P02': ['한옥'], 'P03': ['롯데', '롯캐'], 'P04': ['롯데', 'ㄹㄷ'], 'P05': ['고촌']}
            if any(a in r.raw_text for a in home.get(r.participant_id, [])): place = f'집({r.participant_id})'
            elif '학교' in r.raw_text and r.participant_id == 'P03': place = '김포제일고'
            zone = dict(zip(pl.place_text, pl.zone_id)).get(place, '')
        m = re.search(r'(?<![/\d])([123])(?![/\d])', r.raw_text); inten = int(m.group(1)) if m else None
        ws2.append([r.raw_id, r.participant_id, r.raw_text, r.sent_at, od, ot, prec, place, zone, inten, why, 'N', r.source, r.source_ref,
                    '', '', '', '', 'kakao_export' if r.source == 'kakao' else 'capture_not_attached'])

n = ws2.max_row + 200
for rng_, f in [(f'B2:B{n}', '"P00,P01,P02,P03,P04,P05"'), (f'G2:G{n}', '"exact,hour,part,unknown"'), (f'J2:J{n}', '"1,2,3"'), (f'L2:L{n}', '"Y,N"')]:
    dv = DataValidation(type='list', formula1=f, allow_blank=True); ws2.add_data_validation(dv); dv.add(rng_)
dvp = DataValidation(type='list', formula1=f"=장소_대응표!$A$2:$A${len(pl) + 1}", allow_blank=True); ws2.add_data_validation(dvp); dvp.add(f'H2:H{n}')

W = {'A': 9, 'B': 8, 'C': 46, 'D': 17, 'E': 12, 'F': 8, 'G': 9, 'H': 20, 'I': 7, 'J': 7, 'K': 40, 'L': 8, 'M': 7, 'N': 46, 'O': 20, 'P': 17, 'Q': 17, 'R': 12, 'S': 20}
for k, v in W.items(): ws2.column_dimensions[k].width = v
for k, v in {'A': 22, 'B': 8, 'C': 30, 'D': 26, 'E': 10, 'F': 10}.items(): ws1.column_dimensions[k].width = v
wb.save(OUTF)
print('saved', OUTF, ws2.max_row - 1, 'rows')
